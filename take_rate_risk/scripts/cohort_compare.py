"""
Multi-cohort take-rate / offer-balance / loan-terms / risk comparer.

Phases:
  confirm     — cohort defs + offer/loan counts (gate before full run)
  offers      — offer mix by date, offer times (1/2/3/4+), risk band
  take_rate   — take rate overall + by offer times + by risk band
  loan_terms  — limit / tenor / payment frequency similarity among booked
  risk        — A0,A2,B2,RR1,RR2,GACO,ROA,GM by payweek (A0 >= min_a0)
  all         — confirm tables + offers + take_rate + loan_terms + risk + Excel

Usage (Claw / platform Redshift tool — preferred):
  python cohort_compare.py --config ../configs/my_run.yaml --phase confirm --emit-sql
  # run the emitted .sql via the platform Redshift tool → export CSVs into --data-dir
  python cohort_compare.py --config ../configs/my_run.yaml --phase confirm --data-dir ../output/my_run/raw

Usage (optional local Cred_RS.json fallback):
  python cohort_compare.py --config ../configs/my_run.yaml --phase confirm --use-creds
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
import yaml
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SCRIPT_DIR = Path(__file__).resolve().parent
REPO_ROOT = SCRIPT_DIR.parent.parent
ENGINE_ROOT = SCRIPT_DIR.parent       # .../take_rate_risk
sys.path.insert(0, str(SCRIPT_DIR))
from wiki_metrics import derive_wiki_metrics  # noqa: E402

ROOT = ENGINE_ROOT
DEFAULT_OUT = ROOT / "output"

# CSV names expected when using --data-dir (filled by platform Redshift tool export)
RAW_CONFIRM = "confirm.csv"
RAW_OFFERS = "offers_detail.csv"
RAW_RISK = "risk_raw.csv"

HEADER_FILL = PatternFill("solid", fgColor="0070C0")
HEADER_FONT = Font(bold=True, color="FFFFFF")
METHOD_FILL = PatternFill("solid", fgColor="D9E2F3")
RESULT_FILL = PatternFill("solid", fgColor="E2EFDA")

COLOR_ORDER = ["Green 1", "Green 2", "Blue 1", "Blue 2", "Blue 3", "Yellow", "Other"]
ORDINAL_ORDER = ["1", "2", "3", "4+"]

RISK_DISPLAY = [
    "cohort",
    "payweek",
    "A0_num_Loans",
    "A2_Pct_Loan_DQ",
    "B2_Pct_DQ_Rem_Principal",
    "RR1_Pct_Col_Interest",
    "RR2_Pct_Exp_Interest",
    "GACO",
    "ROA_Cur",
    "ROA_Exp",
    "GM_Cur",
    "GM_Exp",
]

RISK_RENAME = {
    "a0_num_loans": "A0_num_Loans",
    "a2_num_dq": "A2_num_DQ",
    "a1_num_esl": "A1_num_ESL",
    "a3_num_mat": "A3_num_Mat",
    "b0_principal": "B0_Principal",
    "b1_esl_lost_interest": "B1_ESL_Lost_Interest",
    "b2_dq_rem_principal": "B2_DQ_Rem_Principal",
    "b3_dq_rem_interest": "B3_DQ_Rem_Interest",
    "r0_exp_org_interest": "R0_Exp_Org_Interest",
    "r1_asd_interest": "R1_Asd_Interest",
    "r2_asd_lpi": "R2_Asd_LPI",
    "r3_col_interest": "R3_Col_Interest",
    "r4_col_lpi": "R4_Col_LPI",
    "rr3_col_interest": "RR3_Col_Interest",
    "rr4_exp_interest": "RR4_Exp_Interest",
    "dw1_cur_dollarweek": "DW1_Cur_Dollarweek",
    "dw2_exp_dollarweek": "DW2_Exp_Dollarweek",
}


# ---------------------------------------------------------------------------
# Config / DB
# ---------------------------------------------------------------------------


def load_config(path: Path) -> dict[str, Any]:
    cfg = yaml.safe_load(path.read_text(encoding="utf-8"))
    if not isinstance(cfg, dict):
        raise ValueError("Config must be a YAML mapping")
    if "run_name" not in cfg:
        raise ValueError("Config requires run_name")
    cohorts = cfg.get("cohorts") or []
    if len(cohorts) < 2:
        raise ValueError("Need at least 2 cohorts under cohorts:")
    for i, c in enumerate(cohorts):
        if not c.get("name"):
            raise ValueError(f"cohorts[{i}] missing name")
        if not c.get("where") and not c.get("offer_sql"):
            raise ValueError(f"cohort '{c['name']}' needs where: or offer_sql:")
    cfg.setdefault("booking_window_days", 60)
    cfg.setdefault("payweeks", [4, 8, 12, 16, 20, 24])
    cfg.setdefault("min_a0", 30)
    cfg.setdefault("min_loan_amount", 500)
    cfg.setdefault("offer_ordinal_scope", "within_cohort")
    cfg.setdefault("shared_where", [])
    cfg.setdefault("description", "")
    return cfg


def _resolve_creds_path() -> Path:
    env = os.environ.get("CRED_RS_PATH") or os.environ.get("TAKE_RATE_RISK_CREDS")
    if env:
        return Path(env)
    return REPO_ROOT / "Cred_RS.json"


def connect():
    """Optional local fallback only (--use-creds). Prefer platform Redshift tool + --data-dir."""
    try:
        import psycopg2
    except ImportError as e:
        raise ImportError(
            "psycopg2 is only needed for --use-creds. "
            "On Claw, emit SQL and load CSVs via --data-dir instead."
        ) from e

    creds = _resolve_creds_path()
    if not creds.is_file():
        raise FileNotFoundError(
            "Local Cred_RS.json not found (optional fallback only).\n"
            f"  Expected: {creds}\n"
            "  On Claw: use --emit-sql, run SQL with the platform Redshift tool,\n"
            "  export CSVs, then re-run with --data-dir.\n"
            "  Local only: create Cred_RS.json or set CRED_RS_PATH, then --use-creds."
        )
    r = json.loads(creds.read_text(encoding="utf-8"))["redshift"]
    return psycopg2.connect(
        host=r["host"],
        port=r["port"],
        dbname=r["database"],
        user=r["username"],
        password=r["password"],
    )


def _sql_escape_pct(s: str) -> str:
    return s.replace("%", "%%")


def read_sql(sql: str) -> pd.DataFrame:
    with connect() as conn:
        return pd.read_sql(_sql_escape_pct(sql), conn)


def emit_sql_files(cfg: dict, sql_dir: Path, phase: str) -> dict[str, Path]:
    """Write phase SQL for the platform Redshift tool. No DB connection."""
    sql_dir.mkdir(parents=True, exist_ok=True)
    written: dict[str, Path] = {}
    need_confirm = phase in ("confirm", "all")
    need_detail = phase in ("offers", "take_rate", "loan_terms", "all")
    need_risk = phase in ("risk", "all")

    if need_confirm:
        p = sql_dir / "confirm.sql"
        p.write_text(sql_confirm(cfg).strip() + "\n", encoding="utf-8")
        written["confirm"] = p
    if need_detail:
        p = sql_dir / "offers_detail.sql"
        p.write_text(sql_offers_detail(cfg).strip() + "\n", encoding="utf-8")
        written["offers_detail"] = p
    if need_risk:
        p = sql_dir / "risk.sql"
        p.write_text(sql_risk(cfg).strip() + "\n", encoding="utf-8")
        written["risk"] = p

    print(f"\nEmitted SQL under {sql_dir}:")
    for name, path in written.items():
        print(f"  - {name}: {path}")
    print(
        "\nNext (Claw): run each .sql with the platform Redshift tool, "
        f"export CSVs as {RAW_CONFIRM} / {RAW_OFFERS} / {RAW_RISK} into a data dir, "
        "then re-run this script with --data-dir <that dir>."
    )
    return written


def _load_csv(data_dir: Path, name: str) -> pd.DataFrame:
    path = data_dir / name
    if not path.is_file():
        raise FileNotFoundError(
            f"Missing {path}\n"
            "Export the matching Redshift tool result to this CSV, then retry."
        )
    return pd.read_csv(path)


def load_raw_frames(data_dir: Path, phase: str) -> dict[str, pd.DataFrame]:
    """Load CSVs produced by the platform Redshift tool."""
    out: dict[str, pd.DataFrame] = {}
    need_confirm = phase in ("confirm", "all")
    need_detail = phase in ("offers", "take_rate", "loan_terms", "all")
    need_risk = phase in ("risk", "all")
    if need_confirm:
        out["confirm"] = _load_csv(data_dir, RAW_CONFIRM)
    if need_detail:
        out["detail"] = _load_csv(data_dir, RAW_OFFERS)
    if need_risk:
        out["risk_raw"] = _load_csv(data_dir, RAW_RISK)
    return out


def _shared_and(cfg: dict) -> str:
    clauses = cfg.get("shared_where") or []
    if not clauses:
        return "1=1"
    return " AND ".join(f"({c})" for c in clauses)


def _cohort_where(cohort: dict) -> str:
    clauses = cohort.get("where") or []
    if isinstance(clauses, str):
        clauses = [clauses]
    if not clauses:
        return "1=1"
    return " AND ".join(f"({c})" for c in clauses)


def _color_case(alias: str = "o") -> str:
    return f"""CASE
      WHEN {alias}.policy_dcsn_note LIKE '%Green 1%' THEN 'Green 1'
      WHEN {alias}.policy_dcsn_note LIKE '%Green 2%' THEN 'Green 2'
      WHEN {alias}.policy_dcsn_note LIKE '%Blue 1%' THEN 'Blue 1'
      WHEN {alias}.policy_dcsn_note LIKE '%Blue 2%' THEN 'Blue 2'
      WHEN {alias}.policy_dcsn_note LIKE '%Blue 3%' THEN 'Blue 3'
      WHEN {alias}.policy_dcsn_note LIKE '%Yellow%' THEN 'Yellow'
      ELSE 'Other'
    END"""


def _union_offer_ids(cfg: dict) -> str:
    """CTE body: cohort, offer_id from all cohorts."""
    parts = []
    shared = _shared_and(cfg)
    for c in cfg["cohorts"]:
        name = str(c["name"]).replace("'", "''")
        if c.get("offer_sql"):
            body = c["offer_sql"].strip().rstrip(";")
            parts.append(
                f"""SELECT '{name}' AS cohort, x.offer_id
FROM (
{body}
) x"""
            )
        else:
            where = _cohort_where(c)
            parts.append(
                f"""SELECT '{name}' AS cohort, o.offer_id
FROM turbo.luna.acq_offers_info o
WHERE ({shared})
  AND ({where})"""
            )
    return "\nUNION ALL\n".join(parts)


def _partition_keys(cfg: dict) -> str:
    scope = cfg.get("offer_ordinal_scope") or "within_cohort"
    if scope == "across_all":
        return "o.cc_user_id"
    return "c.cohort, o.cc_user_id"


# ---------------------------------------------------------------------------
# SQL
# ---------------------------------------------------------------------------


def sql_confirm(cfg: dict) -> str:
    union = _union_offer_ids(cfg)
    window = int(cfg["booking_window_days"])
    return f"""
WITH cohort_offers AS (
{union}
)
SELECT
  c.cohort,
  COUNT(DISTINCT c.offer_id) AS offers,
  COUNT(DISTINCT CASE
      WHEN bk.loan_id IS NOT NULL
       AND DATEDIFF(day, o.offer_creation_dt, bk.loan_disbursement_dt) < {window}
      THEN bk.loan_id END) AS loans_{window}d,
  COUNT(DISTINCT o.cc_user_id) AS users
FROM cohort_offers c
INNER JOIN turbo.luna.acq_offers_info o ON c.offer_id = o.offer_id
LEFT JOIN turbo.luna.bookings_info_v2 bk ON o.offer_id = bk.luna_offer_id
GROUP BY 1
ORDER BY 1
"""


def sql_offers_detail(cfg: dict) -> str:
    """One row per offer with cohort, risk band, offer ordinal, booking flag."""
    union = _union_offer_ids(cfg)
    window = int(cfg["booking_window_days"])
    part = _partition_keys(cfg)
    color = _color_case("o")
    return f"""
WITH cohort_offers AS (
{union}
),
base AS (
  SELECT
    c.cohort,
    o.offer_id,
    o.cc_user_id,
    o.offer_creation_dt,
    CAST(o.offer_creation_dt AS DATE) AS offer_date,
    o.new_offer_flag,
    {color} AS risk_band,
    CAST(o.offer_conditions.loan_term1 AS INT) AS term1_wks,
    CAST(o.offer_conditions.loan_term2 AS INT) AS term2_wks,
    CAST(o.offer_conditions.loan_term3 AS INT) AS term3_wks,
    CAST(o.offer_conditions.loan_amount1 AS INT) AS loan_amt_max,
    bk.loan_id,
    bk.loan_disbursement_dt,
    bk.loan_disbursed_amt,
    bk.pmt_frequency,
    bk.loan_total_term,
    DATEDIFF(day, o.offer_creation_dt, bk.loan_disbursement_dt) AS days_offer_loan,
    CASE
      WHEN bk.pmt_frequency = 'MONTHLY' THEN bk.loan_total_term * 4.0
      WHEN bk.pmt_frequency = 'WEEKLY' THEN bk.loan_total_term * 1.0
      WHEN bk.pmt_frequency = 'BI_WEEKLY' THEN bk.loan_total_term * 2.0
    END AS loan_tenor_weeks,
    ROW_NUMBER() OVER (
      PARTITION BY {part}
      ORDER BY o.offer_creation_dt, o.offer_id
    ) AS offer_n
  FROM cohort_offers c
  INNER JOIN turbo.luna.acq_offers_info o ON c.offer_id = o.offer_id
  LEFT JOIN turbo.luna.bookings_info_v2 bk ON o.offer_id = bk.luna_offer_id
)
SELECT
  *,
  CASE
    WHEN offer_n = 1 THEN '1'
    WHEN offer_n = 2 THEN '2'
    WHEN offer_n = 3 THEN '3'
    ELSE '4+'
  END AS offer_times,
  CASE
    WHEN loan_id IS NOT NULL AND days_offer_loan IS NOT NULL AND days_offer_loan < {window}
    THEN 1 ELSE 0
  END AS booked_{window}d
FROM base
"""


def sql_risk(cfg: dict) -> str:
    union = _union_offer_ids(cfg)
    window = int(cfg["booking_window_days"])
    payweeks = ", ".join(str(int(p)) for p in cfg["payweeks"])
    return f"""
WITH cohort_offers AS (
{union}
),
offers AS (
  SELECT
    c.cohort,
    o.offer_id,
    bk.loan_id,
    bk.pmt_frequency,
    CASE
      WHEN bk.pmt_frequency = 'MONTHLY' THEN bk.loan_total_term * 4.0
      WHEN bk.pmt_frequency = 'WEEKLY' THEN bk.loan_total_term * 1.0
      WHEN bk.pmt_frequency = 'BI_WEEKLY' THEN bk.loan_total_term * 2.0
    END AS loan_tenor_weeks
  FROM cohort_offers c
  INNER JOIN turbo.luna.acq_offers_info o ON c.offer_id = o.offer_id
  INNER JOIN turbo.luna.bookings_info_v2 bk ON o.offer_id = bk.luna_offer_id
  WHERE bk.loan_id IS NOT NULL
    AND DATEDIFF(day, o.offer_creation_dt, bk.loan_disbursement_dt) < {window}
),
base_dedup AS (
  SELECT *
  FROM (
    SELECT *,
           ROW_NUMBER() OVER (PARTITION BY loan_id ORDER BY offer_id) AS rn
    FROM offers
  )
  WHERE rn = 1
)
SELECT
  b.cohort,
  r.payweek_num AS payweek,
  COUNT(DISTINCT r.loan_id) AS "A0_num_Loans",
  COUNT(DISTINCT CASE WHEN r.dq_flag = 1 THEN r.loan_id END) AS "A2_num_DQ",
  COUNT(DISTINCT CASE WHEN r.esl_flag_mod = 1 THEN r.loan_id END) AS "A1_num_ESL",
  COUNT(DISTINCT CASE WHEN r.maturity_flag_mod = 1 THEN r.loan_id END) AS "A3_num_Mat",
  SUM(r.disbursed_principal_amt) AS "B0_Principal",
  SUM(r.rev_loss_esl_amt) AS "B1_ESL_Lost_Interest",
  SUM(r.dq_flag * r.remaining_principal_amt) AS "B2_DQ_Rem_Principal",
  SUM(r.dq_flag * r.remaining_int_amt) AS "B3_DQ_Rem_Interest",
  SUM(r.expected_original_int_amt) AS "R0_Exp_Org_Interest",
  SUM(r.charged_int_amt) AS "R1_Asd_Interest",
  SUM(r.charged_lpi_amt) AS "R2_Asd_LPI",
  CAST(0 AS FLOAT) AS "R3_Col_Interest",
  CAST(0 AS FLOAT) AS "R4_Col_LPI",
  SUM(r.cum_rev_amt) AS "RR3_Col_Interest",
  SUM(r.expected_rev_amt) AS "RR4_Exp_Interest",
  SUM(
    r.annu_avg_remaining_principal_amt
    * LEAST(r.payweek_num, COALESCE(b.loan_tenor_weeks, r.payweek_num))
  ) AS "DW1_Cur_Dollarweek",
  SUM(
    CASE
      WHEN r.payweek_num >= COALESCE(b.loan_tenor_weeks, r.payweek_num) THEN
        r.annu_avg_remaining_principal_amt * COALESCE(b.loan_tenor_weeks, r.payweek_num)
      ELSE
        r.annu_avg_remaining_principal_amt * r.payweek_num
        + (COALESCE(b.loan_tenor_weeks, r.payweek_num) - r.payweek_num
           + CASE WHEN b.pmt_frequency = 'WEEKLY' THEN 1.0
                  WHEN b.pmt_frequency = 'BI_WEEKLY' THEN 2.0
                  ELSE 4.0 END
          ) / 2.0
        * r.remaining_principal_amt
    END
  ) AS "DW2_Exp_Dollarweek"
FROM base_dedup b
INNER JOIN turbo.luna.risk_user_level_report r ON b.loan_id = r.loan_id
WHERE r.payweek_num IN ({payweeks})
GROUP BY 1, 2
ORDER BY 1, 2
"""


# ---------------------------------------------------------------------------
# Aggregations
# ---------------------------------------------------------------------------


def color_ord(s: pd.Series) -> pd.Series:
    order = {c: i for i, c in enumerate(COLOR_ORDER)}
    return s.map(lambda x: order.get(x, 99))


def ordinal_ord(s: pd.Series) -> pd.Series:
    order = {c: i for i, c in enumerate(ORDINAL_ORDER)}
    return s.map(lambda x: order.get(x, 99))


def booked_mask(df: pd.DataFrame, window: int) -> pd.Series:
    col = f"booked_{window}d"
    if col in df.columns:
        return df[col] == 1
    return (
        df["loan_id"].notna()
        & df["days_offer_loan"].notna()
        & (df["days_offer_loan"] < window)
    )


def offer_summary_by(df: pd.DataFrame, keys: list[str]) -> pd.DataFrame:
    out = (
        df.groupby(keys, dropna=False, as_index=False)
        .agg(offers=("offer_id", "nunique"), users=("cc_user_id", "nunique"))
    )
    totals = df.groupby("cohort", as_index=False).agg(cohort_offers=("offer_id", "nunique"))
    out = out.merge(totals, on="cohort", how="left")
    out["offer_share"] = out["offers"] / out["cohort_offers"].replace(0, np.nan)
    return out


def take_rate_by(df: pd.DataFrame, keys: list[str], window: int) -> pd.DataFrame:
    d = df.copy()
    d["_booked"] = booked_mask(d, window)
    # offer-level: one row per offer
    ol = d.groupby(keys + ["offer_id"], as_index=False).agg(booked=("_booked", "max"))
    out = ol.groupby(keys, as_index=False).agg(
        offers=("offer_id", "nunique"),
        loans=("booked", "sum"),
    )
    out["take_rate"] = out["loans"] / out["offers"].replace(0, np.nan)
    return out


def loan_terms_table(df: pd.DataFrame, window: int, min_amt: float) -> pd.DataFrame:
    d = df.loc[booked_mask(df, window)].copy()
    if d.empty:
        return pd.DataFrame()

    # one loan per offer (already offer×loan); dedupe loan_id
    d = d.sort_values(["loan_id", "offer_id"]).drop_duplicates("loan_id", keep="first")

    d["chose_max_limit"] = (
        d["loan_disbursed_amt"].notna()
        & d["loan_amt_max"].notna()
        & (np.abs(d["loan_disbursed_amt"] - d["loan_amt_max"]) < 1.0)
    )
    d["chose_min_limit"] = (
        d["loan_disbursed_amt"].notna()
        & (np.abs(d["loan_disbursed_amt"] - float(min_amt)) < 1.0)
    )
    # min/max tenor from offered terms (weeks)
    term_cols = [c for c in ("term1_wks", "term2_wks", "term3_wks") if c in d.columns]
    d["offer_min_tenor_wks"] = d[term_cols].min(axis=1, skipna=True)
    d["offer_max_tenor_wks"] = d[term_cols].max(axis=1, skipna=True)
    d["chose_min_tenor"] = (
        d["loan_tenor_weeks"].notna()
        & d["offer_min_tenor_wks"].notna()
        & (np.abs(d["loan_tenor_weeks"] - d["offer_min_tenor_wks"]) < 0.1)
    )
    d["chose_max_tenor"] = (
        d["loan_tenor_weeks"].notna()
        & d["offer_max_tenor_wks"].notna()
        & (np.abs(d["loan_tenor_weeks"] - d["offer_max_tenor_wks"]) < 0.1)
    )
    d["is_weekly"] = d["pmt_frequency"] == "WEEKLY"
    d["is_biweekly"] = d["pmt_frequency"] == "BI_WEEKLY"
    d["is_monthly"] = d["pmt_frequency"] == "MONTHLY"

    return d.groupby("cohort", as_index=False).agg(
        n_loans=("loan_id", "nunique"),
        avg_limit=("loan_amt_max", "mean"),
        avg_disbursed=("loan_disbursed_amt", "mean"),
        pct_choosing_max_limit=("chose_max_limit", "mean"),
        pct_choosing_min_limit=("chose_min_limit", "mean"),
        avg_tenor_weeks=("loan_tenor_weeks", "mean"),
        pct_choosing_min_tenor=("chose_min_tenor", "mean"),
        pct_choosing_max_tenor=("chose_max_tenor", "mean"),
        pct_weekly=("is_weekly", "mean"),
        pct_biweekly=("is_biweekly", "mean"),
        pct_monthly=("is_monthly", "mean"),
    )


def derive_risk(raw: pd.DataFrame, min_a0: int) -> pd.DataFrame:
    if raw is None or raw.empty:
        return pd.DataFrame()
    comp = raw.rename(columns={k: v for k, v in RISK_RENAME.items() if k in raw.columns})
    wiki_cols = [v for v in RISK_RENAME.values() if v in comp.columns]
    for c in wiki_cols + (["payweek"] if "payweek" in comp.columns else []):
        comp[c] = pd.to_numeric(comp[c], errors="coerce")
    g = comp.groupby(["cohort", "payweek"], as_index=False)[wiki_cols].sum()
    g = derive_wiki_metrics(g)
    g = g[g["A0_num_Loans"] >= int(min_a0)].copy()
    cols = [c for c in RISK_DISPLAY if c in g.columns]
    return g[cols].sort_values(["cohort", "payweek"]).reset_index(drop=True)


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------


def style_header(ws, row, n_cols):
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, horizontal="center")


def autosize(ws, max_w=22):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        w = max((len(str(c.value)) if c.value is not None else 0) for c in col) + 2
        ws.column_dimensions[letter].width = min(max_w, max(10, w))


def write_block(ws, start_row, title, method_lines, df, pct_cols=None, num_cols=None, currency_cols=None):
    pct_cols = set(pct_cols or [])
    num_cols = set(num_cols or [])
    currency_cols = set(currency_cols or [])
    r = start_row
    ws.cell(r, 1, title).font = Font(bold=True, size=12)
    r += 1
    ws.cell(r, 1, "METHOD").fill = METHOD_FILL
    ws.cell(r, 1).font = Font(bold=True)
    for i, line in enumerate(method_lines):
        val = str(line)
        if val.startswith("="):
            val = "'" + val
        ws.cell(r + i, 2, val)
    r += max(1, len(method_lines)) + 1
    ws.cell(r, 1, "RESULT").fill = RESULT_FILL
    ws.cell(r, 1).font = Font(bold=True)
    r += 1
    if df is None or df.empty:
        ws.cell(r, 1, "(no rows)")
        return r + 2
    for j, col in enumerate(df.columns, 1):
        ws.cell(r, j, col)
    style_header(ws, r, len(df.columns))
    r += 1
    data_start = r
    for _, row in df.iterrows():
        for j, col in enumerate(df.columns, 1):
            val = row[col]
            if isinstance(val, str) and val.startswith("="):
                val = "'" + val
            cell = ws.cell(r, j, None if pd.isna(val) else val)
            if col in pct_cols and cell.value is not None:
                cell.number_format = "0.00%"
            elif col in currency_cols and cell.value is not None:
                cell.number_format = "$#,##0.00"
            elif col in num_cols and cell.value is not None:
                cell.number_format = "#,##0"
        r += 1
    data_end = r - 1
    for j, col in enumerate(df.columns, 1):
        if col in pct_cols and data_end >= data_start:
            letter = get_column_letter(j)
            ws.conditional_formatting.add(
                f"{letter}{data_start}:{letter}{data_end}",
                ColorScaleRule(
                    start_type="min",
                    start_color="63BE7B",
                    mid_type="percentile",
                    mid_value=50,
                    mid_color="FFEB84",
                    end_type="max",
                    end_color="F8696B",
                ),
            )
    autosize(ws)
    return r + 2


def write_excel(out_xlsx: Path, cfg: dict, tables: dict[str, pd.DataFrame]):
    wb = Workbook()
    window = cfg["booking_window_days"]
    desc = cfg.get("description") or cfg["run_name"]
    cohort_note = " | ".join(
        f"{c['name']}: "
        + (
            "offer_sql"
            if c.get("offer_sql")
            else " AND ".join(c.get("where") or [])
        )
        for c in cfg["cohorts"]
    )
    shared = " AND ".join(cfg.get("shared_where") or []) or "(none)"

    ws = wb.active
    ws.title = "Confirm"
    ws.sheet_view.showGridLines = False
    write_block(
        ws,
        1,
        f"{cfg['run_name']} — Cohort confirmation",
        [
            f"Run: {desc}",
            f"Shared: {shared}",
            f"Cohorts: {cohort_note}",
            f"Loans = booked within {window}d of offer.",
            f"Offer ordinal scope: {cfg.get('offer_ordinal_scope')}",
        ],
        tables.get("confirm", pd.DataFrame()),
        num_cols=[c for c in tables.get("confirm", pd.DataFrame()).columns if c != "cohort"],
    )

    ws2 = wb.create_sheet("Offer_Balance")
    ws2.sheet_view.showGridLines = False
    r = 1
    for key, title in [
        ("offers_overall", "Offers by cohort"),
        ("offers_by_date", "Offers by cohort × offer date"),
        ("offers_by_times", "Offers by cohort × offer times (1/2/3/4+)"),
        ("offers_by_risk", "Offers by cohort × risk band"),
    ]:
        df = tables.get(key, pd.DataFrame())
        r = write_block(
            ws2,
            r,
            title,
            [
                "Goal: check offer mix similarity across cohorts before comparing take rate / risk.",
                "offer_share = offers in cell / total offers in that cohort.",
            ],
            df,
            pct_cols=["offer_share"],
            num_cols=["offers", "users", "cohort_offers"],
        )

    ws3 = wb.create_sheet("Take_Rate")
    ws3.sheet_view.showGridLines = False
    r = 1
    for key, title in [
        ("tr_overall", "Take rate by cohort"),
        ("tr_by_times", "Take rate by cohort × offer times"),
        ("tr_by_risk", "Take rate by cohort × risk band"),
    ]:
        df = tables.get(key, pd.DataFrame())
        r = write_block(
            ws3,
            r,
            title,
            [
                f"Take rate = distinct offers with a loan booked within {window}d / distinct offers.",
            ],
            df,
            pct_cols=["take_rate"],
            num_cols=["offers", "loans"],
        )

    ws4 = wb.create_sheet("Loan_Terms")
    ws4.sheet_view.showGridLines = False
    write_block(
        ws4,
        1,
        "Loan terms similarity (booked loans)",
        [
            f"Among loans booked within {window}d; loan_id deduped.",
            f"pct_choosing_max_limit: |disbursed − offer loan_amount1| < 1.",
            f"pct_choosing_min_limit: |disbursed − {cfg.get('min_loan_amount', 500)}| < 1.",
            "Min/max tenor vs offered term1/2/3 weeks; freq = WEEKLY / BI_WEEKLY / MONTHLY.",
        ],
        tables.get("loan_terms", pd.DataFrame()),
        pct_cols=[
            "pct_choosing_max_limit",
            "pct_choosing_min_limit",
            "pct_choosing_min_tenor",
            "pct_choosing_max_tenor",
            "pct_weekly",
            "pct_biweekly",
            "pct_monthly",
        ],
        num_cols=["n_loans"],
        currency_cols=["avg_limit", "avg_disbursed"],
    )

    ws5 = wb.create_sheet("Risk")
    ws5.sheet_view.showGridLines = False
    write_block(
        ws5,
        1,
        "Risk by cohort × payweek",
        [
            "Source: turbo.luna.risk_user_level_report; formulas via wiki_metrics.py (Luna Wiki).",
            f"Shown only when A0_num_Loans >= {cfg.get('min_a0', 30)}.",
            "A2=DQ loan rate; B2=DQ rem principal/B0; RR1/RR2; GACO; ROA_Cur/Exp; GM_Cur/Exp.",
        ],
        tables.get("risk", pd.DataFrame()),
        pct_cols=[
            "A2_Pct_Loan_DQ",
            "B2_Pct_DQ_Rem_Principal",
            "RR1_Pct_Col_Interest",
            "RR2_Pct_Exp_Interest",
            "GACO",
            "ROA_Cur",
            "ROA_Exp",
            "GM_Cur",
            "GM_Exp",
        ],
        num_cols=["A0_num_Loans", "payweek"],
    )

    out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_xlsx)


# ---------------------------------------------------------------------------
# Run
# ---------------------------------------------------------------------------


def _print_df(title: str, df: pd.DataFrame):
    print(f"\n=== {title} ===")
    if df is None or df.empty:
        print("(no rows)")
        return
    with pd.option_context("display.max_columns", 30, "display.width", 160, "display.float_format", "{:.4f}".format):
        print(df.to_string(index=False))


def run(
    cfg: dict,
    phase: str,
    out_dir: Path,
    *,
    data_dir: Path | None = None,
    use_creds: bool = False,
) -> dict[str, pd.DataFrame]:
    window = int(cfg["booking_window_days"])
    tables: dict[str, pd.DataFrame] = {}

    need_detail = phase in ("offers", "take_rate", "loan_terms", "all")
    need_confirm = phase in ("confirm", "all")
    need_risk = phase in ("risk", "all")

    raw: dict[str, pd.DataFrame] = {}
    if data_dir is not None:
        raw = load_raw_frames(data_dir, phase)
    elif not use_creds:
        raise SystemExit(
            "No data source.\n"
            "  Claw (preferred): --emit-sql → Redshift tool → --data-dir <csv folder>\n"
            "  Local fallback:   --use-creds (requires Cred_RS.json)"
        )

    if need_confirm or phase == "confirm":
        confirm = raw["confirm"] if data_dir is not None else read_sql(sql_confirm(cfg))
        tables["confirm"] = confirm
        _print_df("1) Cohort confirmation (offers / loans / users)", confirm)

    if phase == "confirm":
        return tables

    detail = None
    if need_detail:
        if data_dir is not None:
            print("\nLoading offer-level detail from --data-dir...")
            detail = raw["detail"]
        else:
            print("\nPulling offer-level detail (may take a bit)...")
            detail = read_sql(sql_offers_detail(cfg))
        # sort helpers
        if "risk_band" in detail.columns:
            detail["_c"] = color_ord(detail["risk_band"])
        if "offer_times" in detail.columns:
            detail["_o"] = ordinal_ord(detail["offer_times"])

    if phase in ("offers", "all") and detail is not None:
        overall = (
            detail.groupby("cohort", as_index=False)
            .agg(offers=("offer_id", "nunique"), users=("cc_user_id", "nunique"))
        )
        by_date = offer_summary_by(detail, ["cohort", "offer_date"]).sort_values(
            ["cohort", "offer_date"]
        )
        by_times = offer_summary_by(detail, ["cohort", "offer_times"])
        by_times = (
            by_times.assign(_o=ordinal_ord(by_times["offer_times"]))
            .sort_values(["cohort", "_o"])
            .drop(columns=["_o"])
        )
        by_risk = offer_summary_by(detail, ["cohort", "risk_band"])
        by_risk = (
            by_risk.assign(_c=color_ord(by_risk["risk_band"]))
            .sort_values(["cohort", "_c"])
            .drop(columns=["_c"])
        )
        tables["offers_overall"] = overall
        tables["offers_by_date"] = by_date
        tables["offers_by_times"] = by_times
        tables["offers_by_risk"] = by_risk
        _print_df("2a) Offers by cohort", overall)
        _print_df("2b) Offers by cohort × offer times", by_times)
        _print_df("2c) Offers by cohort × risk band", by_risk)
        _print_df("2d) Offers by cohort × offer date (head)", by_date.head(40))

    if phase in ("take_rate", "all") and detail is not None:
        tr = take_rate_by(detail, ["cohort"], window)
        tr_t = take_rate_by(detail, ["cohort", "offer_times"], window)
        tr_t = tr_t.assign(_o=ordinal_ord(tr_t["offer_times"])).sort_values(["cohort", "_o"]).drop(columns=["_o"])
        tr_r = take_rate_by(detail, ["cohort", "risk_band"], window)
        tr_r = tr_r.assign(_c=color_ord(tr_r["risk_band"])).sort_values(["cohort", "_c"]).drop(columns=["_c"])
        tables["tr_overall"] = tr
        tables["tr_by_times"] = tr_t
        tables["tr_by_risk"] = tr_r
        _print_df("3a) Take rate by cohort", tr)
        _print_df("3b) Take rate by cohort × offer times", tr_t)
        _print_df("3c) Take rate by cohort × risk band", tr_r)

    if phase in ("loan_terms", "all") and detail is not None:
        lt = loan_terms_table(detail, window, cfg.get("min_loan_amount", 500))
        tables["loan_terms"] = lt
        _print_df("4) Loan terms by cohort", lt)

    if need_risk:
        if data_dir is not None:
            print("\nLoading risk aggregates from --data-dir...")
            risk_src = raw["risk_raw"]
        else:
            print("\nPulling risk aggregates...")
            risk_src = read_sql(sql_risk(cfg))
        risk = derive_risk(risk_src, cfg.get("min_a0", 30))
        tables["risk"] = risk
        _print_df(f"5) Risk (A0 >= {cfg.get('min_a0', 30)})", risk)

    if phase == "all":
        out_dir.mkdir(parents=True, exist_ok=True)
        # save CSVs
        for name, df in tables.items():
            if df is not None and not df.empty:
                df.to_csv(out_dir / f"{name}.csv", index=False)
        xlsx = out_dir / f"{cfg['run_name']}_cohort_compare.xlsx"
        write_excel(xlsx, cfg, tables)
        print(f"\nExcel -> {xlsx}")

    return tables


def main():
    ap = argparse.ArgumentParser(description="Multi-cohort take rate / risk comparer")
    ap.add_argument("--config", required=True, help="Path to YAML config")
    ap.add_argument(
        "--phase",
        default="all",
        choices=["confirm", "offers", "take_rate", "loan_terms", "risk", "all"],
    )
    ap.add_argument("--out-dir", default=None, help="Output directory (default output/<run_name>)")
    ap.add_argument(
        "--emit-sql",
        action="store_true",
        help="Write phase SQL files for the platform Redshift tool; do not query",
    )
    ap.add_argument(
        "--data-dir",
        default=None,
        help="Folder with confirm.csv / offers_detail.csv / risk_raw.csv from Redshift tool",
    )
    ap.add_argument(
        "--use-creds",
        action="store_true",
        help="Optional local fallback: query Redshift via Cred_RS.json (not for Claw)",
    )
    args = ap.parse_args()

    cfg_path = Path(args.config)
    cfg = load_config(cfg_path)
    out_dir = Path(args.out_dir) if args.out_dir else DEFAULT_OUT / cfg["run_name"]
    print(f"Run: {cfg['run_name']} | phase={args.phase}")
    print(f"Config: {cfg_path}")

    if args.emit_sql:
        emit_sql_files(cfg, out_dir / "sql", args.phase)
        return

    data_dir = Path(args.data_dir) if args.data_dir else None
    if data_dir is None and not args.use_creds:
        raise SystemExit(
            "Specify --emit-sql, or --data-dir <csvs from Redshift tool>, "
            "or --use-creds for local Cred_RS.json."
        )
    run(cfg, args.phase, out_dir, data_dir=data_dir, use_creds=args.use_creds)


if __name__ == "__main__":
    main()
